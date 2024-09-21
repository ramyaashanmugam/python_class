import openpyxl 
import os

#to create a own/cutom error
class exmp_f_error(Exception):
    def __init__(self, message):
        self.message = message
        super().__init__(self.message)


try:

    r = open("shreena.txt", "r") 
    d=r.read()
    print(d)
    r.close()
    try:
         
        with open("example.txt", "r") as ex1_file:
                print (ex1_file.read())
                ex1_file.close()

    except:
        raise exmp_f_error("example.txt not found") 
    with open("example.txt", "a") as ex2_file:
        ex2_file.write("hello ramya \n")
        ex2_file.write("hello vino ")
    
    print("end line")

    if os.path.exists("example.xslx"):
        print("File exists!")
    else:
        print("File does not exist!")
        with open("example.xslx", "w") as ex_file:
            wb = openpyxl.Workbook() 
            sheet = wb.active 
            c1 = sheet.cell(row = 1, column = 1) 
            c1.value = "s.no"
            c2 = sheet.cell(row = 1, column = 2) 
            c1.value = "names"

except Exception as ex:
    print (ex)
    if "No such file or directory" in str(ex):
        print("please create file before reading")
    if "example.txt not found" in str(ex):
         print("hello")
         with open("example.txt", "w") as ex_file:
            ex_file.write("hello shreena")
